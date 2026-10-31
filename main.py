from fastapi import FastAPI, UploadFile, File, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from docx import Document
from calculations import calculate_rsa
import fitz
import io
import zipfile

app = FastAPI(
    title="YAD Documents Corrector Python Engine",
    version="0.2.0"
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)

MAX_FILE_SIZE = 25 * 1024 * 1024


# -----------------------------
# ROOT AND HEALTH CHECK
# -----------------------------

@app.get("/")
def root():
    return {
        "service": "YAD Python Engine",
        "status": "running",
        "version": app.version
    }


@app.get("/health")
def health():
    return {
        "ok": True,
        "service": "yad-python-engine"
    }


# -----------------------------
# UPLOAD VALIDATION
# -----------------------------

async def read_upload(upload):
    data = await upload.read()

    if not data:
        raise HTTPException(400, "Uploaded file is empty.")

    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(413, "File exceeds 25 MB.")

    return data


# -----------------------------
# DOCUMENT ANALYSIS
# -----------------------------

@app.post("/api/analyze")
async def analyze_document(file: UploadFile = File(...)):
    filename = (file.filename or "").lower()
    data = await read_upload(file)

    try:
        if filename.endswith(".docx"):
            doc = Document(io.BytesIO(data))

            paragraphs = [
                p.text for p in doc.paragraphs
                if p.text.strip()
            ]

            tables = []

            for table in doc.tables:
                for row in table.rows:
                    tables.append([
                        cell.text for cell in row.cells
                    ])

            return {
                "filename": file.filename,
                "file_type": "docx",
                "paragraph_count": len(paragraphs),
                "table_count": len(doc.tables),
                "text": "\n".join(paragraphs),
                "tables": tables
            }

        if filename.endswith(".pdf"):
            pdf = fitz.open(
                stream=data,
                filetype="pdf"
            )

            pages = [
                page.get_text("text")
                for page in pdf
            ]

            return {
                "filename": file.filename,
                "file_type": "pdf",
                "page_count": len(pdf),
                "text": "\n".join(pages)
            }

        raise HTTPException(
            415,
            "Only DOCX and PDF files are supported."
        )

    except HTTPException:
        raise

    except Exception:
        raise HTTPException(
            400,
            "Could not read this document."
        )


# -----------------------------
# ZIP INSPECTION
# -----------------------------

@app.post("/api/zip/check")
async def check_zip(file: UploadFile = File(...)):
    if not (file.filename or "").lower().endswith(".zip"):
        raise HTTPException(415, "Please upload a ZIP file.")

    data = await read_upload(file)

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = [
                item.filename
                for item in archive.infolist()
                if not item.is_dir()
                and item.filename.lower().endswith(
                    (".pdf", ".docx")
                )
            ]

        return {
            "filename": file.filename,
            "document_count": len(names),
            "documents": names
        }

    except zipfile.BadZipFile:
        raise HTTPException(400, "Invalid ZIP file.")


# -----------------------------
# RSA AND MORTGAGE CALCULATIONS
# -----------------------------

class CalculationRequest(BaseModel):
    rsa_balance: str
    dob: str
    calculation_year: int = 2026


@app.post("/api/calculate")
def calculate_documents(request: CalculationRequest):
    try:
        return calculate_rsa(
            balance=request.rsa_balance,
            dob=request.dob,
            calculation_year=request.calculation_year
        )

    except ValueError as error:
        raise HTTPException(
            status_code=400,
            detail=str(error)
        )
class BatchDocument(BaseModel):
    name: str
    base64: str


class BatchAnalysisRequest(BaseModel):
    documents: list[BatchDocument]


@app.post("/api/analyze/batch")
def analyze_batch(request: BatchAnalysisRequest):
    import base64
    import binascii

    if not request.documents:
        raise HTTPException(400, "Upload at least one document.")

    if len(request.documents) > 30:
        raise HTTPException(413, "Maximum 30 documents per batch.")

    results = []
    total_size = 0

    for item in request.documents:
        try:
            raw = base64.b64decode(item.base64, validate=True)
        except (ValueError, binascii.Error):
            raise HTTPException(400, f"Invalid file data: {item.name}")

        total_size += len(raw)

        if not raw:
            raise HTTPException(400, f"Empty file: {item.name}")

        if total_size > 25 * 1024 * 1024:
            raise HTTPException(413, "Combined files exceed 25 MB.")

        try:
            if item.name.lower().endswith(".docx"):
                doc = Document(io.BytesIO(raw))

                paragraphs = [
                    p.text.strip()
                    for p in doc.paragraphs
                    if p.text.strip()
                ]

                tables = [
                    [cell.text.strip() for cell in row.cells]
                    for table in doc.tables
                    for row in table.rows
                ]

                text_content = "\n".join(
                    paragraphs +
                    [" | ".join(row) for row in tables]
                )

                results.append({
                    "filename": item.name,
                    "file_type": "docx",
                    "paragraph_count": len(paragraphs),
                    "table_count": len(doc.tables),
                    "word_count": len(text_content.split()),
                    "text_preview": text_content[:2500],
                    "status": "success"
                })

            elif item.name.lower().endswith(".pdf"):
                pdf = fitz.open(
                    stream=raw,
                    filetype="pdf"
                )

                text_content = "\n".join(
                    page.get_text("text")
                    for page in pdf
                )

                results.append({
                    "filename": item.name,
                    "file_type": "pdf",
                    "page_count": len(pdf),
                    "word_count": len(text_content.split()),
                    "text_preview": text_content[:2500],
                    "status": "success"
                })

            else:
                results.append({
                    "filename": item.name,
                    "status": "unsupported"
                })

        except Exception:
            results.append({
                "filename": item.name,
                "status": "failed",
                "message": "Could not read this document."
            })

    return {
        "service": "yad-python-engine",
        "document_count": len(results),
        "results": results
    }
    
        
# ==========================================
# PYTHON DOCUMENT CORRECTION API
# ==========================================

import base64
import binascii

from corrector import correct_docx


class CorrectionCustomer(BaseModel):
    name: str
    dob: str
    employer_address: str
    rsa_balance: str
    rsa_pin: str
    account_number: str
    calculation_year: int = 2026


class CorrectionDocument(BaseModel):
    name: str
    base64: str


class BatchCorrectionRequest(BaseModel):
    customer: CorrectionCustomer
    documents: list[CorrectionDocument]


@app.post("/api/correct/batch")
def correct_documents_batch(request: BatchCorrectionRequest):

    if not request.documents:
        raise HTTPException(
            400,
            "Upload at least one DOCX document."
        )

    if len(request.documents) > 30:
        raise HTTPException(
            413,
            "Maximum 30 documents per batch."
        )

    try:
        calculations = calculate_rsa(
            balance=request.customer.rsa_balance,
            dob=request.customer.dob,
            calculation_year=request.customer.calculation_year
        )
    except ValueError as error:
        raise HTTPException(400, str(error))

    customer = request.customer.model_dump()

    results = []
    total_size = 0

    for item in request.documents:

        if not item.name.lower().endswith(".docx"):
            results.append({
                "filename": item.name,
                "status": "unsupported",
                "message": "Only DOCX files can be corrected."
            })
            continue

        try:
            raw = base64.b64decode(
                item.base64,
                validate=True
            )
        except (ValueError, binascii.Error):
            raise HTTPException(
                400,
                "Invalid file data: " + item.name
            )

        total_size += len(raw)

        if not raw:
            raise HTTPException(
                400,
                "Empty document: " + item.name
            )

        if total_size > MAX_FILE_SIZE:
            raise HTTPException(
                413,
                "Combined documents exceed 25 MB."
            )

        try:
            corrected_bytes, qc = correct_docx(
                raw,
                customer,
                calculations
            )

            results.append({
                "filename": item.name,
                "status": qc["status"],
                "corrected_base64": base64.b64encode(
                    corrected_bytes
                ).decode("ascii"),
                "quality_check": qc
            })

        except (ValueError, KeyError) as error:
            results.append({
                "filename": item.name,
                "status": "needs_review",
                "message": str(error)
            })

        except Exception:
            results.append({
                "filename": item.name,
                "status": "failed",
                "message": "Unexpected correction error."
            })

    return {
        "service": "yad-python-engine",
        "calculation_engine": "python",
        "document_count": len(results),
        "calculations": calculations,
        "results": results
            }
    
